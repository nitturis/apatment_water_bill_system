#!/usr/bin/env python3
"""Generate water bills with solar-as-common and solar-by-series splits."""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from statistics import mean, median

import openpyxl
from openpyxl.utils import get_column_letter


DEFAULT_TANKER_COST = (74 * 950) + (64 * 740)


@dataclass(frozen=True)
class MeterRow:
    name: str
    previous: float
    current: float

    @property
    def consumption(self) -> float:
        return self.current - self.previous


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Compute monthly water bills from NN_WM_2026.xlsx using two strategies: "
            "solar as common and solar bill per series."
        )
    )
    parser.add_argument(
        "--input",
        default="NN_WM_2026.xlsx",
        help="Input Excel workbook. Default: NN_WM_2026.xlsx",
    )
    parser.add_argument(
        "--sheet",
        default="water_meter_reading_quarter",
        help="Meter reading sheet name. Default: water_meter_reading_quarter",
    )
    parser.add_argument(
        "--previous-column",
        default="March_Reading",
        help="Previous reading column. Default: March_Reading",
    )
    parser.add_argument(
        "--current-column",
        default="Apr_Reading",
        help="Current reading column. Default: Apr_Reading",
    )
    parser.add_argument(
        "--month-label",
        default="April 2026",
        help="Month label for output workbook. Default: April 2026",
    )
    parser.add_argument(
        "--tanker-cost",
        type=float,
        default=DEFAULT_TANKER_COST,
        help=f"Total tanker cost in rupees. Default: {DEFAULT_TANKER_COST:.0f}",
    )
    parser.add_argument(
        "--output",
        default="April_2026_Water_Bill_Two_Strategies.xlsx",
        help="Output Excel workbook path.",
    )
    parser.add_argument(
        "--ignore-solar-series",
        default="",
        help=(
            "Comma-separated solar series to ignore, for faulty meters. "
            "Example: 6,9,10,11,14,18,22"
        ),
    )
    parser.add_argument(
        "--faulty-solar-fill-strategy",
        choices=["no-correction", "mean", "median"],
        default="no-correction",
        help=(
            "How to fill ignored/faulty solar meters. no-correction bills them as 0; "
            "mean/median estimate from working solar series. Default: no-correction."
        ),
    )
    return parser.parse_args()


def column_indexes(header_row: tuple[object, ...]) -> dict[str, int]:
    return {str(value): index for index, value in enumerate(header_row) if value is not None}


def read_meter_rows(
    workbook_path: Path,
    sheet_name: str,
    previous_column: str,
    current_column: str,
) -> list[MeterRow]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Sheet {sheet_name!r} is empty")

    columns = column_indexes(rows[0])
    required = ["Flat No", previous_column, current_column]
    missing = [column for column in required if column not in columns]
    if missing:
        raise ValueError(f"Missing required columns in {sheet_name!r}: {', '.join(missing)}")

    result: list[MeterRow] = []
    for row in rows[1:]:
        flat_name = row[columns["Flat No"]]
        if flat_name is None:
            continue
        previous = row[columns[previous_column]] or 0
        current = row[columns[current_column]] or 0
        result.append(MeterRow(str(flat_name), float(previous), float(current)))

    return result


def series_from_flat(flat_name: str) -> str:
    match = re.search(r"(\d+)$", flat_name)
    if not match:
        return flat_name
    return match.group(1)[-2:]


def normalize_series(series_values: str | list[str] | set[str]) -> set[str]:
    if isinstance(series_values, str):
        values = [value.strip() for value in series_values.split(",")]
    else:
        values = [str(value).strip() for value in series_values]

    normalized: set[str] = set()
    for value in values:
        if not value:
            continue
        match = re.search(r"(\d+)$", value)
        normalized.add(match.group(1)[-2:].zfill(2) if match else value)
    return normalized


def split_rows(rows: list[MeterRow]) -> tuple[list[MeterRow], list[MeterRow], list[MeterRow]]:
    flat_rows: list[MeterRow] = []
    solar_rows: list[MeterRow] = []
    common_rows: list[MeterRow] = []

    for row in rows:
        if row.name.startswith("SOLAR_"):
            solar_rows.append(row)
        elif row.name.startswith("Common_"):
            common_rows.append(row)
        else:
            flat_rows.append(row)

    return flat_rows, solar_rows, common_rows


def solar_series(row: MeterRow) -> str:
    return series_from_flat(row.name.replace("SOLAR_", ""))


def solar_consumption_by_series(
    solar_rows: list[MeterRow],
    ignored_solar_series: set[str],
    fill_strategy: str,
) -> tuple[dict[str, float], dict[str, str], float | None]:
    raw_by_series = {solar_series(row): row.consumption for row in solar_rows}
    working_values = [
        consumption
        for series, consumption in raw_by_series.items()
        if series not in ignored_solar_series
    ]

    fill_value: float | None = None
    if ignored_solar_series and fill_strategy != "no-correction":
        if not working_values:
            raise ValueError("Cannot fill faulty solar meters because no working solar meters are available")
        fill_value = mean(working_values) if fill_strategy == "mean" else median(working_values)

    billed_by_series: dict[str, float] = {}
    status_by_series: dict[str, str] = {}
    for series, consumption in raw_by_series.items():
        if series not in ignored_solar_series:
            billed_by_series[series] = consumption
            status_by_series[series] = "Included"
        elif fill_strategy == "no-correction":
            billed_by_series[series] = 0
            status_by_series[series] = "Ignored - faulty meter"
        else:
            billed_by_series[series] = fill_value or 0
            status_by_series[series] = f"Filled - faulty meter ({fill_strategy})"

    return billed_by_series, status_by_series, fill_value


def money(value: float) -> float:
    return round(value, 2)


def build_records(
    rows: list[MeterRow],
    tanker_cost: float,
    ignored_solar_series: set[str] | None = None,
    faulty_solar_fill_strategy: str = "no-correction",
) -> tuple[dict[str, float], list[dict[str, object]], list[dict[str, object]]]:
    flat_rows, solar_rows, common_rows = split_rows(rows)
    ignored_solar_series = ignored_solar_series or set()
    solar_by_series, solar_status_by_series, solar_fill_value = solar_consumption_by_series(
        solar_rows,
        ignored_solar_series,
        faulty_solar_fill_strategy,
    )

    flat_count = len(flat_rows)
    if flat_count == 0:
        raise ValueError("No flat meter rows found")

    individual_consumption = sum(row.consumption for row in flat_rows)
    ignored_solar_consumption = sum(
        row.consumption
        for row in solar_rows
        if solar_series(row) in ignored_solar_series
    )
    solar_consumption = sum(solar_by_series.values())
    common_consumption = sum(row.consumption for row in common_rows)
    total_consumption = individual_consumption + solar_consumption + common_consumption
    if total_consumption <= 0:
        raise ValueError("Total consumption must be greater than zero")

    unit_cost = tanker_cost / total_consumption
    individual_cost_pool = individual_consumption * unit_cost
    solar_cost_pool = solar_consumption * unit_cost
    common_cost_pool = common_consumption * unit_cost
    solar_common_per_flat = solar_cost_pool / flat_count
    common_per_flat = common_cost_pool / flat_count

    flats_by_series: dict[str, list[str]] = defaultdict(list)
    for row in flat_rows:
        flats_by_series[series_from_flat(row.name)].append(row.name)

    solar_name_by_series = {
        solar_series(row): row.name for row in solar_rows
    }

    flat_records: list[dict[str, object]] = []
    for row in flat_rows:
        series = series_from_flat(row.name)
        series_flat_count = len(flats_by_series[series])
        individual_share = row.consumption * unit_cost
        solar_series_share = (solar_by_series.get(series, 0) * unit_cost) / series_flat_count
        total_solar_common = individual_share + solar_common_per_flat + common_per_flat
        total_solar_series = individual_share + solar_series_share + common_per_flat

        flat_records.append(
            {
                "Flat": row.name,
                "Series": series,
                "Previous Reading": row.previous,
                "Current Reading": row.current,
                "Consumption": row.consumption,
                "Individual Share": individual_share,
                "Solar Share Common": solar_common_per_flat,
                "Solar Share Series": solar_series_share,
                "Common Share": common_per_flat,
                "Total Solar Common": total_solar_common,
                "Total Solar Series": total_solar_series,
                "Difference B-A": total_solar_series - total_solar_common,
            }
        )

    series_records: list[dict[str, object]] = []
    for series in sorted(flats_by_series):
        solar_consumed = solar_by_series.get(series, 0)
        flat_count_in_series = len(flats_by_series[series])
        solar_cost = solar_consumed * unit_cost
        series_records.append(
            {
                "Series": series,
                "Solar Meter": solar_name_by_series.get(series, ""),
                "Solar Status": solar_status_by_series.get(series, "No solar meter"),
                "Flats": flat_count_in_series,
                "Solar Consumption": solar_consumed,
                "Solar Cost Pool": solar_cost,
                "Solar Per Flat": solar_cost / flat_count_in_series,
            }
        )

    summary = {
        "Flat count": flat_count,
        "Total tanker cost": tanker_cost,
        "Individual consumption": individual_consumption,
        "Solar consumption": solar_consumption,
        "Ignored solar consumption": ignored_solar_consumption,
        "Common consumption": common_consumption,
        "Total billable consumption": total_consumption,
        "Unit cost": unit_cost,
        "Individual cost pool": individual_cost_pool,
        "Solar cost pool": solar_cost_pool,
        "Common cost pool": common_cost_pool,
        "Solar as common per flat": solar_common_per_flat,
        "Common facility per flat": common_per_flat,
        "Ignored solar series": ", ".join(sorted(ignored_solar_series)) or "None",
        "Faulty solar fill strategy": faulty_solar_fill_strategy,
        "Faulty solar fill value": solar_fill_value if solar_fill_value is not None else "N/A",
        "Strategy A total recovery": sum(float(row["Total Solar Common"]) for row in flat_records),
        "Strategy B total recovery": sum(float(row["Total Solar Series"]) for row in flat_records),
        "Strategy A min bill": min(float(row["Total Solar Common"]) for row in flat_records),
        "Strategy A max bill": max(float(row["Total Solar Common"]) for row in flat_records),
        "Strategy A average bill": mean(float(row["Total Solar Common"]) for row in flat_records),
        "Strategy B min bill": min(float(row["Total Solar Series"]) for row in flat_records),
        "Strategy B max bill": max(float(row["Total Solar Series"]) for row in flat_records),
        "Strategy B average bill": mean(float(row["Total Solar Series"]) for row in flat_records),
    }

    return summary, flat_records, series_records


def append_table(worksheet, headers: list[str], records: list[dict[str, object]]) -> None:
    worksheet.append(headers)
    for record in records:
        worksheet.append([record[header] for header in headers])


def format_sheet(worksheet) -> None:
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = '#,##0.00'
    for column in range(1, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 20
    worksheet.freeze_panes = "A2"


def write_output(
    output_path: Path,
    month_label: str,
    summary: dict[str, float],
    flat_records: list[dict[str, object]],
    series_records: list[dict[str, object]],
) -> None:
    workbook = openpyxl.Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Metric", "Value"])
    summary_sheet.append(["Month", month_label])
    for key, value in summary.items():
        summary_sheet.append([key, value])

    common_sheet = workbook.create_sheet("Solar As Common")
    append_table(
        common_sheet,
        [
            "Flat",
            "Series",
            "Previous Reading",
            "Current Reading",
            "Consumption",
            "Individual Share",
            "Solar Share Common",
            "Common Share",
            "Total Solar Common",
        ],
        flat_records,
    )

    series_sheet = workbook.create_sheet("Solar By Series")
    append_table(
        series_sheet,
        [
            "Flat",
            "Series",
            "Previous Reading",
            "Current Reading",
            "Consumption",
            "Individual Share",
            "Solar Share Series",
            "Common Share",
            "Total Solar Series",
            "Difference B-A",
        ],
        flat_records,
    )

    series_summary_sheet = workbook.create_sheet("Series Summary")
    append_table(
        series_summary_sheet,
        [
            "Series",
            "Solar Meter",
            "Solar Status",
            "Flats",
            "Solar Consumption",
            "Solar Cost Pool",
            "Solar Per Flat",
        ],
        series_records,
    )

    for worksheet in workbook.worksheets:
        format_sheet(worksheet)

    workbook.save(output_path)


def print_summary(output_path: Path, summary: dict[str, float]) -> None:
    print(f"Created: {output_path.resolve()}")
    print(f"Flat count: {int(summary['Flat count'])}")
    print(f"Total tanker cost: Rs {money(summary['Total tanker cost']):,.2f}")
    print(f"Total billable consumption: {summary['Total billable consumption']:,.2f}")
    print(f"Unit cost: Rs {summary['Unit cost']:,.4f}")
    print(f"Solar as common per flat: Rs {summary['Solar as common per flat']:,.2f}")
    print(f"Common facility per flat: Rs {summary['Common facility per flat']:,.2f}")
    print(
        "Strategy A bill range: "
        f"Rs {money(summary['Strategy A min bill']):,.2f} - "
        f"Rs {money(summary['Strategy A max bill']):,.2f}"
    )
    print(
        "Strategy B bill range: "
        f"Rs {money(summary['Strategy B min bill']):,.2f} - "
        f"Rs {money(summary['Strategy B max bill']):,.2f}"
    )


def main() -> None:
    args = parse_args()
    input_path = Path(args.input)
    output_path = Path(args.output)

    rows = read_meter_rows(
        input_path,
        args.sheet,
        args.previous_column,
        args.current_column,
    )
    ignored_solar_series = normalize_series(args.ignore_solar_series)
    summary, flat_records, series_records = build_records(
        rows,
        args.tanker_cost,
        ignored_solar_series,
        args.faulty_solar_fill_strategy,
    )
    write_output(output_path, args.month_label, summary, flat_records, series_records)
    print_summary(output_path, summary)


if __name__ == "__main__":
    main()
